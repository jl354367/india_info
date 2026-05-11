import re
from io import BytesIO
from typing import List, Dict, Optional

import pandas as pd
import openpyxl

from app.models.report_item import ReportItem
import logging

SHEET_NAME = "India Index"
REQUIRED_COLUMNS = ["Rank", "Sector", "Topic"]

EXCEL_COLUMNS = [
    "Rank",
    "Sector",
    "Topic",
    "Website_Link",
    "Timeline_Years",
    "Before_vs_After",
    "Trend",
    "Key_Stats",
    "Detailed_Summary",
]

COLUMN_MAP = {
    "Rank": "Rank",
    "Sector": "Sector",
    "Topic": "Topic",
    "Website_Link": "WebsiteLink",
    "Timeline_Years": "TimelineYears",
    "Before_vs_After": "BeforeVsAfter",
    "Trend": "Trend",
    "Key_Stats": "KeyStats",
    "Detailed_Summary": "DetailedSummary",
}

_HYPERLINK_FORMULA_RE = re.compile(r'HYPERLINK\(\s*"([^"]+)"', re.IGNORECASE)


def _as_bytes(file_bytes) -> bytes:
    if hasattr(file_bytes, "read"):
        data = file_bytes.read()
        try:
            file_bytes.seek(0)
        except Exception:
            pass
        return data
    return file_bytes


def _load_website_links(raw: bytes) -> Dict[int, Optional[str]]:
    """Return {data_row_index (0-based): url} for the Website_Link column.

    Prefers the cell's hyperlink target (so display text like "Open" still resolves
    to the underlying URL), then a =HYPERLINK("url",...) formula, then the raw
    string value of the cell.
    """
    wb = openpyxl.load_workbook(BytesIO(raw), data_only=False)
    if SHEET_NAME not in wb.sheetnames:
        return {}
    ws = wb[SHEET_NAME]

    website_col = None
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if isinstance(cell.value, str) and cell.value.strip() == "Website_Link":
            website_col = cell.column
            break
    if website_col is None:
        return {}

    out: Dict[int, Optional[str]] = {}
    for i, row in enumerate(
        ws.iter_rows(min_row=2, min_col=website_col, max_col=website_col)
    ):
        cell = row[0]
        url: Optional[str] = None
        if cell.hyperlink and getattr(cell.hyperlink, "target", None):
            url = cell.hyperlink.target
        elif isinstance(cell.value, str):
            m = _HYPERLINK_FORMULA_RE.search(cell.value)
            url = m.group(1) if m else cell.value
        if isinstance(url, str):
            url = url.strip() or None
        out[i] = url
    return out


def parse_excel(file_bytes) -> List[ReportItem]:
    raw = _as_bytes(file_bytes)

    df = pd.read_excel(BytesIO(raw), sheet_name=SHEET_NAME, engine="openpyxl")
    df.columns = [str(c).strip() for c in df.columns]

    missing = [c for c in EXCEL_COLUMNS if c not in df.columns]
    if missing:
        logging.error(
            f"Missing expected columns in excel: {missing}. Found: {list(df.columns)}"
        )
        raise ValueError(f"Excel is missing required columns: {missing}")

    df = df[EXCEL_COLUMNS].reset_index(drop=True)
    for col in EXCEL_COLUMNS:
        df[col] = df[col].map(lambda x: x.strip() if isinstance(x, str) else x)

    website_links = _load_website_links(raw)

    items = []
    for idx, row in df.iterrows():
        data = {
            COLUMN_MAP[col]: (row[col] if pd.notnull(row[col]) else None)
            for col in EXCEL_COLUMNS
        }
        # Override Website_Link with the hyperlink target if openpyxl found one.
        hyperlink_url = website_links.get(int(idx))
        if hyperlink_url:
            data["WebsiteLink"] = hyperlink_url

        if any(not data.get(COLUMN_MAP[col]) for col in REQUIRED_COLUMNS):
            logging.warning(f"Skipping row with missing required fields: {data}")
            continue
        try:
            items.append(ReportItem(**data))
        except Exception as e:
            logging.error(f"Error parsing row: {data} | {e}")
    return items
